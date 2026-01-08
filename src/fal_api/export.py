"""
Layer Export Utilities.
분해된 레이어를 다양한 형식(PNG, PSD, PPTX, ZIP)으로 내보냅니다.
"""

import os
import zipfile
from pathlib import Path
from typing import List, Optional
from PIL import Image
from pptx import Presentation
from psd_tools import PSDImage


class LayerExporter:
    """
    분해된 레이어를 다양한 형식으로 내보내는 클래스.
    지원 형식: PNG, PSD, PPTX, ZIP
    """

    def __init__(self, layer_files: List[str]):
        """
        Args:
            layer_files: 레이어 이미지 파일 경로 리스트
        """
        self.layer_files = layer_files
        self._validate_files()

    def _validate_files(self):
        """파일 존재 여부 확인"""
        for f in self.layer_files:
            if not os.path.exists(f):
                raise FileNotFoundError(f"파일을 찾을 수 없습니다: {f}")

    def _get_image_size(self) -> tuple:
        """첫 번째 이미지의 크기를 반환"""
        with Image.open(self.layer_files[0]) as img:
            return img.size

    @staticmethod
    def _px_to_emu(px: int, dpi: int = 96) -> int:
        """픽셀을 EMU(English Metric Units)로 변환"""
        inch = px / dpi
        emu = inch * 914400
        return int(emu)

    def to_pptx(self, output_path: str) -> str:
        """
        레이어들을 PPTX 파일로 내보냅니다.
        모든 레이어가 하나의 슬라이드에 겹쳐져 배치됩니다.

        Args:
            output_path: 출력 PPTX 파일 경로

        Returns:
            저장된 파일 경로
        """
        width, height = self._get_image_size()

        prs = Presentation()
        prs.slide_width = self._px_to_emu(width)
        prs.slide_height = self._px_to_emu(height)

        # 빈 슬라이드 추가
        slide = prs.slides.add_slide(prs.slide_layouts[6])

        # 모든 레이어를 순서대로 추가
        for img_path in self.layer_files:
            slide.shapes.add_picture(
                img_path,
                left=0,
                top=0,
                width=self._px_to_emu(width),
                height=self._px_to_emu(height)
            )

        prs.save(output_path)
        print(f"PPTX 저장됨: {output_path}")
        return output_path

    def to_psd(self, output_path: str) -> str:
        """
        레이어들을 PSD(Photoshop) 파일로 내보냅니다.

        Args:
            output_path: 출력 PSD 파일 경로

        Returns:
            저장된 파일 경로
        """
        layers = []
        for path in self.layer_files:
            img = Image.open(path).convert('RGBA')
            layers.append(img)

        width, height = layers[0].size
        psd = PSDImage.new(mode='RGBA', size=(width, height))

        for i, img in enumerate(layers):
            name = f"Layer {i + 1}"
            layer = psd.create_pixel_layer(image=img, name=name)
            psd.append(layer)

        psd.save(output_path)
        print(f"PSD 저장됨: {output_path}")
        return output_path

    def to_zip(self, output_path: str) -> str:
        """
        레이어들을 ZIP 파일로 묶어서 내보냅니다.

        Args:
            output_path: 출력 ZIP 파일 경로

        Returns:
            저장된 파일 경로
        """
        with zipfile.ZipFile(output_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
            for i, img_path in enumerate(self.layer_files):
                arcname = f"layer_{i+1}.png"
                zipf.write(img_path, arcname)

        print(f"ZIP 저장됨: {output_path}")
        return output_path

    def export_all(self, output_dir: str, base_name: str = "layers") -> dict:
        """
        모든 형식(PPTX, PSD, ZIP)으로 한 번에 내보냅니다.

        Args:
            output_dir: 출력 디렉토리
            base_name: 파일명 기본값

        Returns:
            각 형식별 저장된 파일 경로 딕셔너리
        """
        output_path = Path(output_dir)
        output_path.mkdir(parents=True, exist_ok=True)

        results = {
            "pptx": self.to_pptx(str(output_path / f"{base_name}.pptx")),
            "psd": self.to_psd(str(output_path / f"{base_name}.psd")),
            "zip": self.to_zip(str(output_path / f"{base_name}.zip")),
            "png_files": self.layer_files.copy(),
        }

        print(f"\n모든 형식 내보내기 완료!")
        print(f"  - PPTX: {results['pptx']}")
        print(f"  - PSD: {results['psd']}")
        print(f"  - ZIP: {results['zip']}")
        print(f"  - PNG: {len(results['png_files'])}개 파일")

        return results


def export_layers(
    layer_files: List[str],
    output_dir: str,
    formats: Optional[List[str]] = None
) -> dict:
    """
    레이어 파일들을 지정된 형식으로 내보내는 간편 함수.

    Args:
        layer_files: 레이어 이미지 파일 경로 리스트
        output_dir: 출력 디렉토리
        formats: 내보낼 형식 리스트 ['pptx', 'psd', 'zip']. None이면 모두 내보냄.

    Returns:
        각 형식별 저장된 파일 경로 딕셔너리
    """
    exporter = LayerExporter(layer_files)

    if formats is None:
        return exporter.export_all(output_dir)

    output_path = Path(output_dir)
    output_path.mkdir(parents=True, exist_ok=True)

    results = {"png_files": layer_files.copy()}

    if 'pptx' in formats:
        results['pptx'] = exporter.to_pptx(str(output_path / "layers.pptx"))
    if 'psd' in formats:
        results['psd'] = exporter.to_psd(str(output_path / "layers.psd"))
    if 'zip' in formats:
        results['zip'] = exporter.to_zip(str(output_path / "layers.zip"))

    return results
